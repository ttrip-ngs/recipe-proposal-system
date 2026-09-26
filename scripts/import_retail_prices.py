"""小売物価統計調査の都市別小売価格から食材価格表 YAML を生成するスクリプト.

総務省 小売物価統計調査(動向編)「主要品目の都市別小売価格」の食料品部分
(銘柄符号 1001〜1999) を e-Stat から Excel で取得し, 指定都市の価格を
src/recipe_system/cost/data/retail_prices.yaml に書き出す.

使い方 (月 1 回, 新しい月の公表後に実行し, 差分をレビューして取り込む):
    1. e-Stat の統計表一覧で対象月の「1001 うるち米 … 2183 学校給食」の表の
       statInfId を確認する
       https://www.e-stat.go.jp/stat-search/files?page=1&layout=datalist&toukei=00200571&tstat=000000680001&cycle=1
    2. uv run python scripts/import_retail_prices.py --stat-inf-id 000040506177

ファイル取得にアプリケーション ID は不要 (e-Stat の公開ファイルダウンロードを使う).
"""

from __future__ import annotations

import argparse
import io
import re
import sys
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx
import yaml
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = REPO_ROOT / "src" / "recipe_system" / "cost" / "data" / "retail_prices.yaml"
DOWNLOAD_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={id}&fileKind=0"

# 食料品 (外食・酒類を除く) の銘柄符号範囲
FOOD_CODE_MIN = 1001
FOOD_CODE_MAX = 1999

_PERIOD_RE = re.compile(r"(\d{4})年(\d{1,2})月")


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="都市別小売価格から食材価格表 YAML を生成")
    parser.add_argument("--stat-inf-id", required=True, help="e-Stat の統計表 ID (statInfId)")
    parser.add_argument("--area", default="福岡市", help="対象都市 (表の列見出しと一致させる)")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def download(stat_inf_id: str) -> bytes:
    url = DOWNLOAD_URL.format(id=stat_inf_id)
    resp = httpx.get(url, follow_redirects=True, timeout=60.0)
    resp.raise_for_status()
    return resp.content


def parse_workbook(content: bytes, area: str) -> tuple[str, dict[int, dict[str, Any]]]:
    """Excel から (対象月 YYYY-MM, {銘柄符号: {name, unit, price_jpy}}) を取り出す."""
    wb = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))

    header_idx = next(
        (i for i, r in enumerate(rows) if "銘柄符号" in r),
        None,
    )
    if header_idx is None:
        raise ValueError("見出し行 (銘柄符号) が見つかりません")
    header = rows[header_idx]
    col_code = header.index("銘柄符号")
    col_name = header.index("品目")
    col_unit = header.index("単位")
    col_period = header.index("時間軸\uff08月\uff09")  # 全角括弧の見出し
    if area not in header:
        raise ValueError(f"都市 {area} の列が見つかりません")
    col_area = header.index(area)

    period: str | None = None
    items: dict[int, dict[str, Any]] = {}
    for r in rows[header_idx + 1 :]:
        code = r[col_code]
        if not isinstance(code, int) or not FOOD_CODE_MIN <= code <= FOOD_CODE_MAX:
            continue
        period_cell = r[col_period]
        if period is None and isinstance(period_cell, str):
            m = _PERIOD_RE.search(period_cell)
            if m:
                period = f"{m.group(1)}-{int(m.group(2)):02d}"
        price = r[col_area]
        # "..." (調査なし) や "-" (出回りなし) は価格なしとして除外する
        if not isinstance(price, int | float) or price <= 0:
            continue
        items[code] = {"name": str(r[col_name]), "unit": str(r[col_unit]), "price_jpy": price}

    if period is None:
        raise ValueError("対象月を特定できません")
    return period, items


def render(*, period: str, area: str, stat_inf_id: str, items: dict[int, dict[str, Any]]) -> str:
    header = (
        "# 自動生成ファイル. 手で編集しない (scripts/import_retail_prices.py で再生成する).\n"
        "# 出典: 総務省 小売物価統計調査(動向編) 主要品目の都市別小売価格\n"
    )
    doc = {
        "source": "小売物価統計調査(動向編) 主要品目の都市別小売価格",
        "source_url": DOWNLOAD_URL.format(id=stat_inf_id),
        "area": area,
        "period": period,
        "generated_at": datetime.now(tz=UTC).date().isoformat(),
        "items": dict(sorted(items.items())),
    }
    return header + yaml.safe_dump(doc, allow_unicode=True, sort_keys=False)


def main() -> int:
    args = _parse_args()
    content = download(args.stat_inf_id)
    period, items = parse_workbook(content, args.area)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        render(period=period, area=args.area, stat_inf_id=args.stat_inf_id, items=items),
        encoding="utf-8",
    )
    print(f"{args.area} {period}: {len(items)} 品目を {args.output} に書き出しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
